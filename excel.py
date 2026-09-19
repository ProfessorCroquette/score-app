"""Read/write .xlsx files. No DB, no menu logic."""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---------- input ----------

def read_students(path: str) -> tuple[list[str], list[dict]]:
    """
    Read an .xlsx file with this layout:
        | NIS     | Name    | Quiz | Midterm | Final |
        | 2024001 | Alice   | 80   | 90      | 70    |

    Returns:
        components: ["Quiz", "Midterm", "Final"]
        students:   [
            {"nis": "2024001", "name": "Alice",
             "Quiz": 80.0, "Midterm": 90.0, "Final": 70.0},
            ...
        ]
    """
    wb = load_workbook(path, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    # first two columns: NIS, Name. The rest are components.
    components = [h for h in headers[2:] if h]

    students = []
    for row in rows[1:]:
        if not row or row[0] is None:
            continue

        nis = str(row[0]).strip()
        name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""

        if not nis and not name:
            continue

        record = {"nis": nis, "name": name}
        for i, comp in enumerate(components, start=2):
            val = row[i] if i < len(row) else None
            try:
                record[comp] = float(val) if val is not None else 0.0
            except (ValueError, TypeError):
                record[comp] = 0.0
        students.append(record)

    return components, students


# ---------- output ----------

def write_results(path: str,
                  components: list[str],
                  results: list[dict],
                  summ: dict,
                  scheme_name: str = "simple") -> None:
    """Write a formatted 3-sheet workbook: Results, Summary, Distribution."""

    wb = Workbook()

    # shared styles
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ---------- Sheet 1: Results ----------
    ws = wb.active
    ws.title = "Results"

    header = ["Rank", "NIS", "Name"] + components + ["Weighted", "Grade"]
    ws.append(header)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for r in results:
        row = [r["rank"], r.get("nis", ""), r["name"]]
        row += [r["scores"].get(c, 0) for c in components]
        row += [round(r["weighted"], 2), r["grade"]]
        ws.append(row)

    # borders + centering on data rows
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                            min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = border
            cell.alignment = center

    # left-align the name column
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=3):
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 22
    for i in range(4, 4 + len(components) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 10

    # freeze the header row so it stays visible while scrolling
    ws.freeze_panes = "A2"

    # ---------- Sheet 2: Summary ----------
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    ws2.append(["Students",       summ["total"]])
    ws2.append(["Class average",  round(summ["average"], 2)])
    ws2.append(["Passing",        summ["passing"]])
    ws2.append(["Failing",        summ["total"] - summ["passing"]])
    ws2.append(["Pass rate",      f"{summ['pass_rate'] * 100:.1f}%"])
    ws2.append(["Grading scheme", scheme_name])

    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 15

    # ---------- Sheet 3: Distribution ----------
    ws3 = wb.create_sheet("Distribution")
    ws3.append(["Grade", "Count"])
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    # get the grades list from calc.SCHEMES (lazy import to avoid cycle)
    from calc import SCHEMES
    scheme = SCHEMES.get(scheme_name, SCHEMES["simple"])
    grade_list = scheme["grades"]

    counts = {g: 0 for g in grade_list}
    for r in results:
        counts[r["grade"]] = counts.get(r["grade"], 0) + 1

    for g in grade_list:
        ws3.append([g, counts.get(g, 0)])

    ws3.column_dimensions["A"].width = 10
    ws3.column_dimensions["B"].width = 10

    wb.save(path)
